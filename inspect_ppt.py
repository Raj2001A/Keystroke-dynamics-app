import collections
import collections.abc
from pptx import Presentation

ppt_path = r"R:\Project papers\Type2Branch\repo\Type2Branch_Final_Viva.pptx"
prs = Presentation(ppt_path)

print(f"Total slides: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            texts.append(shape.text.replace("\n", " ").replace("\v", " ")[:150])
    print(" | ".join(texts))
